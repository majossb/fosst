import csv
import io
from datetime import datetime
from openpyxl import Workbook, load_workbook
from django.db import transaction
from apps.capacitaciones.models import Trabajador
from apps.accounts.models import TipoDocumento

HEADERS = [
    "nombre",
    "apellidos",
    "tipo_documento",
    "documento",
    "email",
    "telefono",
    "tipo_vinculacion",
    "tipo_contrato",
    "fecha_ingreso",
    "salario",
]

def generar_plantilla_excel() -> bytes:
    """Genera una plantilla Excel (.xlsx) limpia con encabezados y filas de ejemplo."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Plantilla Trabajadores"

    ws.append(HEADERS)

    # Filas de ejemplo
    ws.append([
        "Juan Carlos",
        "Pérez Gómez",
        "CC",
        "1012345678",
        "juan.perez@ejemplo.com",
        "3001234567",
        "dependiente",
        "indefinido",
        "2026-01-15",
        "2500000.00",
    ])
    ws.append([
        "María Fernanda",
        "López Ruiz",
        "CC",
        "1098765432",
        "maria.lopez@ejemplo.com",
        "3109876543",
        "independiente",
        "prestacion_servicios",
        "2026-02-01",
        "3000000.00",
    ])

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()


def procesar_importacion_trabajadores(file_obj, empresa, user) -> dict:
    """
    Procesa importación masiva de trabajadores desde Excel (.xlsx) o CSV (.csv).
    Valida fila por fila dentro de una transacción individual por registro.
    """
    filename = getattr(file_obj, "name", "").lower()
    filas = []

    if filename.endswith(".csv"):
        content = file_obj.read()
        if isinstance(content, bytes):
            content = content.decode("utf-8-sig", errors="ignore")
        reader = csv.DictReader(io.StringIO(content))
        for idx, row in enumerate(reader, start=2):
            filas.append((idx, {k.strip().lower(): str(v).strip() for k, v in row.items() if k}))
    else:
        wb = load_workbook(file_obj, data_only=True)
        ws = wb.active
        header_names = [str(cell.value).strip().lower() for cell in ws[1] if cell.value]

        for idx, row_cells in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if not any(row_cells):
                continue
            row_dict = {}
            for col_idx, val in enumerate(row_cells):
                if col_idx < len(header_names):
                    row_dict[header_names[col_idx]] = str(val).strip() if val is not None else ""
            filas.append((idx, row_dict))

    total_filas = len(filas)
    creadas = 0
    fallidas = 0
    reporte = []

    for num_fila, data in filas:
        doc = data.get("documento", "")
        nombre = data.get("nombre", "")
        apellidos = data.get("apellidos", "")

        if not doc or not nombre:
            fallidas += 1
            reporte.append({
                "fila": num_fila,
                "estado": "error",
                "documento": doc,
                "mensaje": "Nombre y número de documento son obligatorios."
            })
            continue

        # Verificar duplicados en la empresa
        if Trabajador.objects.filter(empresa=empresa, documento=doc).exists():
            fallidas += 1
            reporte.append({
                "fila": num_fila,
                "estado": "error",
                "documento": doc,
                "mensaje": f"Ya existe un trabajador registrado con el documento {doc} en la empresa."
            })
            continue

        tipo_doc = data.get("tipo_documento", "CC").upper()
        if tipo_doc not in [choice[0] for choice in TipoDocumento.choices]:
            tipo_doc = TipoDocumento.CC

        tipo_vinc = data.get("tipo_vinculacion", "dependiente").lower()
        if tipo_vinc not in [choice[0] for choice in Trabajador.TipoVinculacion.choices]:
            tipo_vinc = Trabajador.TipoVinculacion.DEPENDIENTE

        tipo_cont = data.get("tipo_contrato", "indefinido").lower()
        if tipo_cont not in [choice[0] for choice in Trabajador.TipoContrato.choices]:
            tipo_cont = Trabajador.TipoContrato.INDEFINIDO

        # Parsear fecha
        raw_fecha = data.get("fecha_ingreso", "")
        fecha_ingreso = None
        if raw_fecha:
            for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%Y/%m/%d", "%d-%m-%Y"):
                try:
                    fecha_ingreso = datetime.strptime(raw_fecha[:10], fmt).date()
                    break
                except ValueError:
                    pass

        # Parsear salario
        raw_salario = data.get("salario", "0").replace(",", ".").replace("$", "").strip()
        try:
            salario = float(raw_salario) if raw_salario else 0.0
        except ValueError:
            salario = 0.0

        try:
            with transaction.atomic():
                Trabajador.objects.create(
                    empresa=empresa,
                    nombre=f"{nombre} {apellidos}".strip(),
                    documento=doc,
                    tipo_vinculacion=tipo_vinc,
                    tipo_contrato=tipo_cont,
                    fecha_ingreso=fecha_ingreso,
                    remuneracion_monto=salario,
                )
            creadas += 1
            reporte.append({
                "fila": num_fila,
                "estado": "exito",
                "documento": doc,
                "mensaje": "Trabajador creado exitosamente."
            })
        except Exception as e:
            fallidas += 1
            reporte.append({
                "fila": num_fila,
                "estado": "error",
                "documento": doc,
                "mensaje": f"Error al guardar registro: {str(e)}"
            })

    return {
        "total_filas": total_filas,
        "creadas": creadas,
        "fallidas": fallidas,
        "reporte": reporte,
    }
